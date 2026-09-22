from pathlib import Path
import csv
from pptx import Presentation
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent.parent
DASHBOARD_DIR = BASE / "dashboard" / "final"
ANALYSIS_DIR = BASE / "analysis"
INSIGHTS_FILE = ANALYSIS_DIR / "insights" / "final_insights.csv"
KPI_FILE = ANALYSIS_DIR / "kpi" / "final_kpi_summary.csv"
RECOMMENDATIONS_FILE = BASE / "recommendations" / "final_recommendations.csv"
DASHBOARD_IMAGE = DASHBOARD_DIR / "week12_final_kpi_dashboard.png"

OUTPUT_FILE = Path(__file__).resolve().parent / "CadetX_Week12_Final_Presentation.pptx"


def read_csv(path):
    if not path.exists():
        return []
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


kpis = read_csv(KPI_FILE)
insights = read_csv(INSIGHTS_FILE)
recommendations = read_csv(RECOMMENDATIONS_FILE)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.35), Inches(12), Inches(0.7)
    )
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True

    if subtitle:
        box2 = slide.shapes.add_textbox(
            Inches(0.6), Inches(1.05), Inches(12), Inches(0.45)
        )
        p2 = box2.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(14)


def add_bullets(slide, items, left=0.8, top=1.6, width=11.8, height=5.2):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(item)
        p.font.size = Pt(18)
        p.space_after = Pt(12)
        p.level = 0


# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
title = slide.shapes.add_textbox(
    Inches(0.8), Inches(2.1), Inches(11.8), Inches(1.2)
)
p = title.text_frame.paragraphs[0]
p.text = "CadetX Warehouse Analytics"
p.font.size = Pt(34)
p.font.bold = True

sub = slide.shapes.add_textbox(
    Inches(0.8), Inches(3.35), Inches(11.8), Inches(1)
)
p = sub.text_frame.paragraphs[0]
p.text = "Week 12 – Final Analytics, Insights & Recommendations"
p.font.size = Pt(22)

# Slide 2: Project Overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Project Overview")
add_bullets(slide, [
    "End-to-end warehouse analytics project covering product, inventory, sales, supplier and customer data.",
    "Week 12 consolidates final KPIs, business insights, dashboard outputs and recommendations.",
    "Python and Pandas were used for analytics and output generation.",
    "Final deliverables include KPI summary, insights, recommendations and dashboard."
])

# Slide 3: KPI Snapshot
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Final KPI Snapshot")

items = []
for row in kpis:
    kpi = row.get("KPI", "")
    value = row.get("Value", "")
    unit = row.get("Unit", "")
    if kpi:
        items.append(f"{kpi}: {value} {unit}")

if not items:
    items = [
        "Total Products: 30 Count",
        "Total Inventory Records: 180 Count",
        "Total Sales Order Lines: 130402 Count",
        "Total Suppliers: 8 Count",
        "Total Branches: 6 Count",
        "Total Customers: 500 Count",
        "Total Units Sold: 1368534 Units",
        "Total Sales Value: 25447392560 Currency",
        "Total Current Stock: 19303266 Units"
    ]

add_bullets(slide, items, top=1.4, height=5.8)

# Slide 4: Business Insights
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Key Business Insights")

items = []
for row in insights:
    insight = row.get("Insight", "")
    kpi = row.get("KPI", "")
    value = row.get("Value", "")
    if insight:
        items.append(f"{kpi} ({value}): {insight}")

if not items:
    items = [
        "The business works with 8 suppliers, supporting procurement and supply operations.",
        "The analysis covers 6 branches, providing a multi-branch operational view.",
        "The dataset contains 500 customers.",
        "Total units sold are 1368534 Units, representing overall sales volume.",
        "Total sales value recorded is 25447392560 Currency.",
        "Current inventory stands at 19303266 Units."
    ]

add_bullets(slide, items, top=1.4, height=5.8)

# Slide 5: Recommendations
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Business Recommendations")

items = []
for row in recommendations:
    kpi = row.get("KPI", "")
    recommendation = (
        row.get("Recommendation")
        or row.get("Recommendation / Action")
        or row.get("Action")
        or ""
    )
    if recommendation:
        items.append(f"{kpi}: {recommendation}")

if not items:
    items = [
        "Use customer-level sales patterns to support segmentation and business decisions.",
        "Track unit sales trends regularly to identify demand patterns and improve replenishment planning.",
        "Monitor sales value trends and connect revenue trends with product, customer and branch performance.",
        "Monitor current inventory against demand to reduce excess stock and potential stockout risk."
    ]

add_bullets(slide, items, top=1.4, height=5.8)

# Slide 6: Final Dashboard
if DASHBOARD_IMAGE.exists():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Final KPI Dashboard")
    slide.shapes.add_picture(
        str(DASHBOARD_IMAGE),
        Inches(0.7),
        Inches(1.25),
        width=Inches(11.9)
    )

# Slide 7: Conclusion
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Conclusion")
add_bullets(slide, [
    "The Week 12 phase consolidates the warehouse analytics project into final decision-support outputs.",
    "Final KPIs provide a concise view of products, inventory, customers, suppliers, branches and sales.",
    "Insights translate the analytical results into understandable business observations.",
    "Recommendations provide practical areas for monitoring inventory, sales and customer activity.",
    "The final dashboard brings the major KPIs together for management-level review."
])

prs.save(OUTPUT_FILE)

print("=" * 55)
print("FINAL PRESENTATION CREATED SUCCESSFULLY")
print("=" * 55)
print(f"Presentation: {OUTPUT_FILE}")
print(f"Slides: {len(prs.slides)}")
